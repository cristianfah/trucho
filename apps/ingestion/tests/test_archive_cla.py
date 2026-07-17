"""Tests del importer del dataset de Internet Archive (Excel construido en el test)."""

from pathlib import Path

import openpyxl
import pytest

from trucho_ingest.models import AwardTier, SourceConfidence
from trucho_ingest.scrapers.archive_cla import parse_workbook, tier_from_award

HEADER_CLASSIC = [
    "Award", "Category", "Brand", "Product", "Title", "Agency", "City",
    "Production", "City", "Country", "Source", "Comments",
]
HEADER_WITH_ENTRANT = [
    "Award", "Category", "Entrant", "Brand", "Product", "Title", "Agency",
    "City", "Production", "City", "Country", "Source", "Comments",
]


@pytest.fixture(scope="module")
def xlsx(tmp_path_factory) -> Path:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    ws54 = wb.create_sheet("1954")
    ws54.append(HEADER_CLASSIC)
    ws54.append(["Grand Prix", "Grand Prix d’ISAS", "Chlorodont Toothpaste", None,
                 "Il Circo", None, None, "Ferry Mayer", None, "Italy", "AW54/LTW", None])
    ws54.append(["Runner-Up", None, "D.H. Evans", None, "Window Display", None, None,
                 "Screenspace Ltd.", None, "United Kingdom", "AW54", None])
    # Award vacío: hereda "Runner-Up" de la fila anterior
    ws54.append([None, None, "Lux Toilet Soap", None, "A Star", None, None,
                 "Pearl & Dean", None, "United Kingdom", "AW54", "live action"])
    # premio a persona/agencia, sin pieza: se descarta
    ws54.append(["Palme d’Or", None, None, None, None, None, None,
                 "Nino Pagot", None, "Italy", "AW54", None])

    ws99 = wb.create_sheet("1999")
    ws99.append(HEADER_WITH_ENTRANT)
    ws99.append(["Grand Prix", None, None, "The Independent", None, "Litany",
                 "Lowe Howard-Spink", "London", "Helen Langridge", "London",
                 "United Kingdom", "IAF99/LTW", None])
    ws99.append(["Agency of the Year", None, None, None, None, None,
                 "DM9DDB", None, None, None, "Brazil", "IAF99", None])
    ws99.append(["Gold", "Beverages", None, "Guinness", "Draught", "Surfer",
                 "AMV BBDO", "London", None, None, "United Kingdom", "IAF99", None])

    ws75 = wb.create_sheet("1975")
    ws75.append(HEADER_CLASSIC)
    ws75.append(["Diploma?", "Food", "Dr.Pepper", None, "Executive Launch", "Y&R",
                 None, None, None, "United States", "AW75", "uncertain record"])

    wb.create_sheet("Unconfirmed").append(HEADER_CLASSIC)
    wb.create_sheet("Indeterminate Year").append(HEADER_CLASSIC)

    path = tmp_path_factory.mktemp("ia") / "cla.xlsx"
    wb.save(path)
    return path


@pytest.fixture(scope="module")
def campaigns(xlsx):
    return list(parse_workbook(xlsx))


def test_parses_piece_rows_only(campaigns):
    titles = {c.title for c in campaigns}
    # Palme d'Or (sin pieza) y Agency of the Year quedan afuera
    assert titles == {"Il Circo", "Window Display", "A Star", "Litany", "Surfer",
                      "Executive Launch"}


def test_award_inherited_from_previous_row(campaigns):
    lux = next(c for c in campaigns if c.title == "A Star")
    assert lux.awards[0].tier == AwardTier.SILVER  # hereda "Runner-Up"
    assert lux.year == 1954


def test_header_variant_with_entrant_column(campaigns):
    litany = next(c for c in campaigns if c.title == "Litany")
    assert litany.brand == "The Independent"
    assert litany.country == "GB"
    assert litany.awards[0].tier == AwardTier.GRAND_PRIX
    assert litany.awards[0].category == "Film"
    assert [a.name for a in litany.agencies][0] == "Lowe Howard-Spink"


def test_category_prefixed_with_film(campaigns):
    surfer = next(c for c in campaigns if c.title == "Surfer")
    assert surfer.awards[0].category == "Film — Beverages"
    assert surfer.awards[0].tier == AwardTier.GOLD


def test_low_confidence_for_1999_2000(campaigns):
    litany = next(c for c in campaigns if c.title == "Litany")
    assert litany.confidence == SourceConfidence.LOW  # error de fechas conocido
    circo = next(c for c in campaigns if c.title == "Il Circo")
    assert circo.confidence == SourceConfidence.NORMAL


def test_low_confidence_for_dubious_award(campaigns):
    dubious = next(c for c in campaigns if c.title == "Executive Launch")
    assert dubious.confidence == SourceConfidence.LOW  # "Diploma?"
    assert dubious.awards[0].tier == AwardTier.SHORTLIST


def test_year_filter(xlsx):
    only_99 = list(parse_workbook(xlsx, year_filter=1999))
    assert {c.year for c in only_99} == {1999}


def test_raw_text_carries_provenance(campaigns):
    litany = next(c for c in campaigns if c.title == "Litany")
    assert "Cannes Lions 1999" in litany.raw_text
    assert "IAF99" in litany.raw_text
    assert litany.source_site == "archive-org-cla"


@pytest.mark.parametrize(
    ("award", "tier"),
    [
        ("Grand Prix", AwardTier.GRAND_PRIX),
        ("Gold", AwardTier.GOLD),
        ("Gold Medal", AwardTier.GOLD),
        ("1st", AwardTier.GOLD),
        ("Silver", AwardTier.SILVER),
        ("SIlver", AwardTier.SILVER),
        ("2nd", AwardTier.SILVER),
        ("Runner-up", AwardTier.SILVER),
        ("Bronze", AwardTier.BRONZE),
        ("3rd", AwardTier.BRONZE),
        ("Diploma", AwardTier.SHORTLIST),
        ("Diploma of Recognition", AwardTier.SHORTLIST),
        ("Honourable Mention", AwardTier.SHORTLIST),
        ("Shortlist", AwardTier.SHORTLIST),
        ("5", AwardTier.SHORTLIST),
        ("Special Jury Prize", AwardTier.WINNER),
        ("Coppa di Venezia", AwardTier.WINNER),
    ],
)
def test_tier_mapping(award, tier):
    assert tier_from_award(award) == tier
